# -*- coding: utf-8 -*-
"""Dung 02 mon KIT: to gap A4 (HTML in) + file Word 06 prompt."""
import sys, io, os, html
sys.path.insert(0, r"C:\Users\vietdt\AppData\Local\Temp\claude\C--Users-vietdt-OneDrive---PetroVietnam-Exploration-Production-Corporation--PVEP--Desktop-PVEP-Zix-CLAUDE\0bd8e559-eb9c-4719-9059-4f34c41ac51d\scratchpad")
from prompts_data import *

OUT = r"C:\Users\vietdt\OneDrive - PetroVietnam Exploration Production Corporation (PVEP)\Desktop\PVEP\Zix\CLAUDE\HTML\SLIDE\AI-SLIDE-WORKSHOP\VP\kit"
os.makedirs(OUT, exist_ok=True)
E = html.escape

# ================================================================ TO GAP A4
def prompt_block(p):
    fields = "".join(f'<div class="f"><b>{E(k)}:</b> {E(v)}</div>' for k, v in p["fields"])
    rules = "".join(f"<li>{E(r)}</li>" for r in p["rules"])
    return f'''<div class="pr">
  <div class="pr__h"><span class="pr__n">{p["n"]}</span><span class="pr__t">{E(p["title"])}</span><span class="pr__tag">{E(p["tag"])}</span></div>
  <div class="pr__run"><b>Chạy ở đâu:</b> {E(p["run"])}</div>
  <div class="pr__body">{fields}<div class="f"><b>Yêu cầu:</b></div><ul>{rules}</ul></div>
</div>'''

six = "".join(f'<div class="six"><span class="six__n">{i+1}</span><div><b>{E(k)}</b><small>{E(d)}</small><em>{E(ex)}</em></div></div>' for i, (k, d, ex) in enumerate(SIX_PARTS))
chk = "".join(f'<div class="ck"><span class="ck__box"></span><div><b>{E(k)}</b><small>{E(d)}</small></div></div>' for k, d in CHECKLIST)
route = "".join(f'<div class="rt"><div class="rt__w">{E(w)}</div><div class="rt__t">{E(t)}</div><div class="rt__n">{E(n)}</div></div>' for w, t, n in TOOL_ROUTE)
mx_head = "".join(f"<th>{E(h)}</th>" for h in MATRIX_HEAD)
mx_rows = "".join("<tr><td class='mx__t'>" + E(t) + "</td>" + "".join(f"<td>{HARVEY[v]}</td>" for v in vals) + "</tr>" for t, vals in MATRIX)
zones = "".join(
    f'<div class="zn" style="--c:{c}"><div class="zn__h"><span>{E(code)}</span>{E(t)}</div><div class="zn__b">' +
    "".join(f"<i>{E(x)}</i>" for x in chips) + f'</div><div class="zn__r">{E(r)}</div></div>'
    for code, t, c, chips, r in ZONES)
leaks = "".join(f"<span>✕ {E(x)}</span>" for x in LEAKS)

page = f'''<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>AI cho Văn phòng PVEP - tờ gập cầm tay</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Sora:wght@600;700;800&family=Be+Vietnam+Pro:ital,wght@0,400;0,600;0,700;1,400&display=swap" rel="stylesheet">
<style>
:root{{--g:#006838;--g2:#00A651;--navy:#0A2A4A;--ink:#16203A;--ink2:#4B5563;--mute:#94A3B8;--line:#D9E2EC;--bg:#F3F6F9}}
*{{box-sizing:border-box;margin:0;padding:0}}
html,body{{background:#cfd8e3;font-family:"Be Vietnam Pro",Arial,Helvetica,sans-serif;color:var(--ink);font-size:8.2pt;line-height:1.32;-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.toolbar{{position:sticky;top:0;z-index:9;background:var(--navy);color:#fff;padding:8px 14px;display:flex;gap:12px;align-items:center;font-size:10.5pt}}
.toolbar b{{font-family:Sora,sans-serif}}
.toolbar button{{margin-left:auto;background:var(--g2);color:#fff;border:0;border-radius:999px;padding:7px 16px;font-weight:700;cursor:pointer;font-family:inherit}}
.toolbar small{{opacity:.8}}
.sheet{{width:297mm;height:210mm;background:#fff;margin:10mm auto;display:grid;grid-template-columns:1fr 1fr;position:relative;box-shadow:0 8px 30px rgba(0,0,0,.18);overflow:hidden;page-break-after:always}}
.sheet:after{{content:"";position:absolute;left:50%;top:0;bottom:0;border-left:1px dashed #b7c3d0}}
.panel{{padding:8mm 8mm 7mm;display:flex;flex-direction:column;gap:2.4mm;min-width:0;overflow:hidden;position:relative}}
.panel__foot{{margin-top:auto;font-size:6.6pt;color:var(--mute);display:flex;justify-content:space-between;border-top:1px solid var(--line);padding-top:1.6mm}}
h1{{font-family:Sora,sans-serif;font-size:15pt;color:var(--g);line-height:1.15}}
h2{{font-family:Sora,sans-serif;font-size:8.6pt;color:var(--navy);text-transform:uppercase;letter-spacing:.06em;border-left:3px solid var(--g2);padding-left:2mm;margin-top:1mm}}
.brand{{display:flex;align-items:center;gap:3mm;border-bottom:2px solid var(--g);padding-bottom:2mm}}
.brand img{{height:11mm;width:auto}}
.brand small{{display:block;color:var(--ink2);font-size:7pt}}
.lead{{font-size:7.8pt;color:var(--ink2)}}
.lead b{{color:var(--g)}}
/* mat 1 */
.sixwrap{{display:grid;grid-template-columns:1fr 1fr;gap:1.6mm}}
.six{{display:flex;gap:1.6mm;border:1px solid var(--line);border-radius:2mm;padding:1.6mm 2mm;background:#fff}}
.six__n{{flex:none;width:5mm;height:5mm;border-radius:50%;background:var(--g);color:#fff;font-family:Sora,sans-serif;font-weight:800;font-size:7.5pt;display:flex;align-items:center;justify-content:center}}
.six b{{display:block;font-size:8pt;color:var(--navy)}}
.six small{{display:block;font-size:6.9pt;color:var(--ink2);line-height:1.25}}
.six em{{display:block;font-size:6.8pt;color:var(--g);margin-top:.6mm;font-style:italic}}
.formula{{background:var(--bg);border-radius:2mm;padding:1.8mm 2.4mm;font-size:7.4pt;font-weight:700;color:var(--navy);text-align:center}}
.ck{{display:flex;gap:2mm;align-items:flex-start;padding:1.2mm 0;border-bottom:1px dotted var(--line)}}
.ck:last-child{{border-bottom:0}}
.ck__box{{flex:none;width:4.2mm;height:4.2mm;border:1.5px solid var(--g);border-radius:1mm;margin-top:.4mm}}
.ck b{{display:block;font-size:7.9pt;color:var(--ink)}}
.ck small{{display:block;font-size:6.9pt;color:var(--ink2)}}
.rule{{background:var(--navy);color:#fff;border-radius:2mm;padding:2mm 2.6mm;font-size:7.6pt;line-height:1.35}}
.rule b{{color:#86efac}}
/* mat 2-3: prompt */
.pr{{border:1px solid var(--line);border-radius:2mm;overflow:hidden;background:#fff}}
.pr__h{{display:flex;align-items:center;gap:2mm;background:var(--g);color:#fff;padding:1.2mm 2.2mm}}
.pr__n{{font-family:Sora,sans-serif;font-weight:800;font-size:9pt;background:rgba(255,255,255,.18);border-radius:1.2mm;padding:0 1.4mm}}
.pr__t{{font-family:Sora,sans-serif;font-weight:700;font-size:8.2pt}}
.pr__tag{{margin-left:auto;font-size:6.4pt;opacity:.9}}
.pr__run{{background:#E6F7EE;color:#064E3B;font-size:6.9pt;padding:1.1mm 2.2mm;border-bottom:1px solid #C6EBD6;line-height:1.3}}
.pr__body{{padding:1.4mm 2.2mm 1.6mm;font-size:7pt;line-height:1.3}}
.pr__body .f b{{color:var(--navy)}}
.pr__body ul{{list-style:none;margin-top:.4mm}}
.pr__body li{{position:relative;padding-left:2.6mm;color:var(--ink2)}}
.pr__body li:before{{content:"›";position:absolute;left:.4mm;color:var(--g2);font-weight:800}}
/* mat 4 */
.rt{{display:grid;grid-template-columns:1.35fr auto 1.1fr;gap:1.6mm;align-items:center;border:1px solid var(--line);border-radius:2mm;padding:1.3mm 2mm;font-size:7pt}}
.rt__w{{color:var(--ink);font-weight:600}}
.rt__t{{font-family:Sora,sans-serif;font-weight:800;color:#fff;background:var(--navy);border-radius:999px;padding:.6mm 2.2mm;font-size:7pt;white-space:nowrap}}
.rt__n{{color:var(--ink2);font-size:6.6pt}}
table.mx{{width:100%;border-collapse:collapse;font-size:6.6pt}}
table.mx th{{font-weight:600;color:var(--ink2);text-align:center;padding:.6mm .4mm;border-bottom:1px solid var(--line);line-height:1.15}}
table.mx td{{text-align:center;padding:.5mm .4mm;font-size:10pt;color:var(--g);line-height:1;border-bottom:1px dotted var(--line)}}
table.mx td.mx__t{{text-align:left;font-size:7.2pt;font-weight:700;color:var(--navy);font-family:Sora,sans-serif}}
.mxl{{font-size:6.3pt;color:var(--ink2);text-align:right}}
.zones{{display:grid;grid-template-columns:repeat(3,1fr);gap:1.6mm}}
.zn{{border:1px solid var(--c);border-radius:2mm;overflow:hidden;display:flex;flex-direction:column}}
.zn__h{{background:var(--c);color:#fff;font-weight:700;font-size:7pt;padding:1mm 1.6mm;display:flex;gap:1.2mm;align-items:center}}
.zn__h span{{background:rgba(255,255,255,.22);border-radius:1mm;padding:0 1.2mm;font-family:Sora,sans-serif;font-size:6.4pt}}
.zn__b{{padding:1.2mm 1.4mm;display:flex;flex-wrap:wrap;gap:.9mm}}
.zn__b i{{font-style:normal;font-size:6.2pt;border:1px solid var(--line);border-radius:999px;padding:.3mm 1.4mm;color:var(--ink);background:#fff;line-height:1.25}}
.zn__r{{margin-top:auto;font-size:6.2pt;color:var(--c);padding:0 1.4mm 1.2mm;font-weight:600;line-height:1.25}}
.leaks{{display:flex;flex-wrap:wrap;gap:1.2mm;font-size:6.6pt;font-weight:700;color:#B91C1C}}
.leaks span{{border:1px solid #FCA5A5;background:#FEF2F2;border-radius:999px;padding:.5mm 1.8mm}}
.print-note{{max-width:297mm;margin:0 auto 12mm;font-size:9.5pt;color:#334155;background:#fff;border-radius:8px;padding:10px 14px}}
@page{{size:A4 landscape;margin:0}}
@media print{{
  html,body{{background:#fff}}
  .toolbar,.print-note{{display:none}}
  .sheet{{margin:0;box-shadow:none;width:297mm;height:210mm}}
}}
</style>
</head>
<body>
<div class="toolbar"><b>Tờ gập cầm tay · AI cho Văn phòng PVEP</b><small>In A4 ngang, 2 mặt (lật cạnh dài), gập đôi. Trang 1 là mặt ngoài, trang 2 là mặt trong.</small><button onclick="window.print()">In / Lưu PDF</button></div>

<!-- ===== TRANG 1 (mặt ngoài): trái = mặt 4 (bìa sau), phải = mặt 1 (bìa trước) ===== -->
<div class="sheet">
  <div class="panel">
    <h2>Việc nào mở công cụ nào</h2>
    {route}
    <h2>Con nào mạnh việc gì</h2>
    <table class="mx"><thead><tr><th style="text-align:left">Công cụ</th>{mx_head}</tr></thead><tbody>{mx_rows}</tbody></table>
    <div class="mxl">● mạnh nhất · ◕ · ◑ · ◔ · ○ không có. Chọn theo việc, không chọn theo thói quen.</div>
    <h2>Trước khi dán vào AI: tài liệu vùng nào</h2>
    <div class="zones">{zones}</div>
    <div class="leaks">{leaks}</div>
    <div class="panel__foot"><span>Mặt 4 · Chọn công cụ và bảo mật dữ liệu</span><span>Lỡ đưa nhầm: báo ngay người phụ trách</span></div>
  </div>
  <div class="panel">
    <div class="brand"><img src="../PVEP.png" alt="PVEP"><div><b style="font-family:Sora,sans-serif;color:var(--g);font-size:9.5pt">Văn phòng PVEP</b><small>Buổi chia sẻ nội bộ · AI cho Văn phòng</small></div></div>
    <h1>AI cho Văn phòng<br>tờ gập cầm tay</h1>
    <p class="lead">1Office giữ quy trình, số, hạn và luồng ký. Tờ này chỉ lo phần <b>nội dung trước khi đưa vào 1Office</b>: hỏi AI cho đúng, kiểm cho kỹ, đưa dữ liệu cho đúng chỗ.</p>
    <h2>Khuôn prompt 06 phần</h2>
    <div class="sixwrap">{six}</div>
    <div class="formula">Vai trò + Bối cảnh/nguồn + Nhiệm vụ + Đối tượng nhận + Đầu ra + Ràng buộc</div>
    <h2>Checklist 30 giây trước khi gửi</h2>
    <div>{chk}</div>
    <div class="rule">AI làm <b>bản nháp và xử lý thông tin</b>. Con người chịu trách nhiệm về <b>quyết định, tính đúng đắn, bảo mật và việc phát hành</b>.</div>
    <div class="panel__foot"><span>Mặt 1 · Khuôn prompt và checklist</span><span>Ban Kiểm soát nội bộ PVEP · 09/2026</span></div>
  </div>
</div>

<!-- ===== TRANG 2 (mặt trong): trái = mặt 2, phải = mặt 3 ===== -->
<div class="sheet">
  <div class="panel">
    <h2>06 prompt Văn phòng · 01-03</h2>
    {prompt_block(PROMPTS[0])}
    {prompt_block(PROMPTS[1])}
    {prompt_block(PROMPTS[2])}
    <div class="panel__foot"><span>Mặt 2 · Công văn đến, văn bản trả lời, chuẩn bị họp</span><span>Bản Word để bôi chép: file 06-prompt-Van-phong.docx</span></div>
  </div>
  <div class="panel">
    <h2>06 prompt Văn phòng · 04-06</h2>
    {prompt_block(PROMPTS[3])}
    {prompt_block(PROMPTS[4])}
    {prompt_block(PROMPTS[5])}
    <div class="panel__foot"><span>Mặt 3 · Biên bản, đón đoàn, truyền thông</span><span>Chỗ không rõ: ghi “Chưa xác định”, không tự suy đoán</span></div>
  </div>
</div>

<div class="print-note"><b>Cách in:</b> Ctrl+P, khổ A4 ngang, lề 0, bật “Background graphics”, in 2 mặt lật theo cạnh dài. Gập đôi theo vạch đứt: mặt có logo là bìa trước.</div>
</body>
</html>
'''
io.open(os.path.join(OUT, "to-gap-A4.html"), "w", encoding="utf-8", newline="\n").write(page)

# ================================================================ FILE WORD 06 PROMPT
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

def para(text="", bold=False, size=None, color=None, italic=False, align=None, space_after=4):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

para("VĂN PHÒNG PVEP · BUỔI CHIA SẺ NỘI BỘ AI CHO VĂN PHÒNG", bold=True, size=9, color="4B5563")
para("06 PROMPT VĂN PHÒNG DÙNG HẰNG TUẦN", bold=True, size=16, color="006838", space_after=2)
para("Bản Word để bôi chép. Mỗi prompt: đọc dòng “Chạy ở đâu”, chép khung xám dán vào AI, đính kèm tài liệu, thay phần trong ngoặc vuông [ ]. "
     "1Office giữ quy trình, số, hạn và luồng ký; prompt chỉ lo phần nội dung trước khi đưa vào 1Office.", size=10, color="4B5563", space_after=8)

for p in PROMPTS:
    para(f"Prompt {p['n']} · {p['title']}", bold=True, size=13, color="0A2A4A", space_after=1)
    para(f"{p['tag']}", italic=True, size=9.5, color="4B5563", space_after=2)
    q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(4)
    r = q.add_run("Chạy ở đâu: "); r.bold = True; r.font.color.rgb = RGBColor.from_string("047857"); r.font.size = Pt(10)
    r = q.add_run(p["run"]); r.font.color.rgb = RGBColor.from_string("064E3B"); r.font.size = Pt(10)
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; shade(c, "F3F6F9")
    c.paragraphs[0].text = ""
    first = True
    for line in prompt_plain(p).split("\n"):
        pp = c.paragraphs[0] if first else c.add_paragraph(); first = False
        pp.paragraph_format.space_after = Pt(1)
        if ":" in line and not line.startswith("-"):
            k, v = line.split(":", 1)
            rk = pp.add_run(k + ":"); rk.bold = True; rk.font.size = Pt(10)
            rv = pp.add_run(v); rv.font.size = Pt(10)
        else:
            rr = pp.add_run(line); rr.font.size = Pt(10)
    para("", space_after=8)

para("Checklist 30 giây trước khi gửi", bold=True, size=13, color="0A2A4A", space_after=2)
for k, d in CHECKLIST:
    pp = doc.add_paragraph(style="List Bullet"); pp.paragraph_format.space_after = Pt(1)
    r = pp.add_run(k + ": "); r.bold = True
    pp.add_run(d)
para("", space_after=4)
para("AI làm bản nháp và xử lý thông tin. Con người chịu trách nhiệm về quyết định, tính đúng đắn, bảo mật và việc phát hành.", italic=True, size=10, color="4B5563")
para("Ban Kiểm soát nội bộ PVEP · 09/2026", size=9, color="94A3B8", align=WD_ALIGN_PARAGRAPH.RIGHT)

doc.save(os.path.join(OUT, "06-prompt-Van-phong.docx"))
print("OK", os.listdir(OUT))
